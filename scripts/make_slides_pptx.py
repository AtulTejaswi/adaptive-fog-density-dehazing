"""Generate the FogSense presentation deck (PPTX) from committed results.

Run:  python scripts/make_slides_pptx.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "figures"
OUT = ROOT / "report" / "FogSense_slides.pptx"

NAVY = RGBColor(0x0B, 0x1F, 0x33)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
AMBER = RGBColor(0xD9, 0x82, 0x2B)
CORAL = RGBColor(0xC8, 0x55, 0x3D)
GRAY = RGBColor(0x6B, 0x7A, 0x8A)
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def rgb(r, g, b):
    return RGBColor(r, g, b)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def _fill(shape, color, transparency=None):
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = color
    if transparency is not None:
        fill.transparency = transparency


def _line(shape, color, width_pt=1.0):
    line = shape.line
    line.color.rgb = color
    line.width = Pt(width_pt)


def add_footer(slide, n):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(6.95),
                                 Inches(12.23), Inches(0.01))
    _fill(line, RGBColor(0xCD, 0xD6, 0xDE))
    line.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(7.02), Inches(6.0), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "FogSense \u00b7 IPA course project"
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY
    tb2 = slide.shapes.add_textbox(Inches(11.9), Inches(7.02), Inches(1.0), Inches(0.3))
    p2 = tb2.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = str(n)
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = GRAY


def content_slide(title, subtitle=None, n=None):
    slide = prs.slides.add_slide(BLANK)
    # left accent
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                                Inches(0.18), Inches(7.5))
    _fill(bar, TEAL)
    bar.line.fill.background()
    # title
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.55))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r.font.name = "Segoe UI"
    if subtitle:
        stb = slide.shapes.add_textbox(Inches(0.55), Inches(0.88), Inches(12.2), Inches(0.3))
        sp = stb.text_frame.paragraphs[0]
        sr = sp.add_run()
        sr.text = subtitle
        sr.font.size = Pt(12.5)
        sr.font.color.rgb = GRAY
        sr.font.name = "Calibri"
    # rule
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.18),
                                 Inches(12.23), Inches(0.015))
    _fill(rule, RGBColor(0xCD, 0xD6, 0xDE))
    rule.line.fill.background()
    if n is not None:
        add_footer(slide, n)
    return slide


def title_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                               Inches(13.333), Inches(7.5))
    _fill(bg, NAVY)
    bg.line.fill.background()
    # teal accent
    acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(1.05),
                                Inches(0.12), Inches(5.4))
    _fill(acc, TEAL)
    acc.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(1.05), Inches(1.0), Inches(7.0), Inches(1.3))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "FogSense"
    r.font.size = Pt(62)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = "Segoe UI"
    sub = slide.shapes.add_textbox(Inches(1.05), Inches(2.25), Inches(7.0), Inches(0.9))
    sp = sub.text_frame.paragraphs[0]
    sr = sp.add_run()
    sr.text = "An Intelligent Multi-Algorithm Framework for Adaptive Haze Removal"
    sr.font.size = Pt(20)
    sr.font.color.rgb = RGBColor(0xB7, 0xD9, 0xDD)
    sr.font.name = "Segoe UI"
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.05), Inches(3.35),
                                 Inches(5.4), Inches(0.03))
    _fill(line, TEAL)
    line.line.fill.background()
    meta = slide.shapes.add_textbox(Inches(1.05), Inches(3.55), Inches(7.0), Inches(1.6))
    mp = meta.text_frame.paragraphs[0]
    mr = mp.add_run()
    mr.text = "IPA Course Project \u00b7 Image Processing & Analysis"
    mr.font.size = Pt(15)
    mr.font.color.rgb = WHITE
    mr2 = mp.add_run()
    mr2.text = "\nClassical image processing only \u00b7 no deep dehazing networks"
    mr2.font.size = Pt(12)
    mr2.font.color.rgb = RGBColor(0xA9, 0xB8, 0xC8)
    mr2.font.name = "Calibri"
    foot = slide.shapes.add_textbox(Inches(1.05), Inches(6.35), Inches(7.0), Inches(0.7))
    fp = foot.text_frame.paragraphs[0]
    fr = fp.add_run()
    fr.text = "October 2026  \u00b7  github.com/AtulTejaswi/adaptive-fog-density-dehazing"
    fr.font.size = Pt(10.5)
    fr.font.color.rgb = RGBColor(0x7F, 0x93, 0xA8)
    # right image card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.05), Inches(0.95),
                                 Inches(4.75), Inches(5.55))
    _fill(card, WHITE)
    card.line.color.rgb = RGBColor(0x2A, 0x4A, 0x66)
    card.line.width = Pt(1.2)
    try:
        slide.shapes.add_picture(str(FIG / "slide_teaser.png"), Inches(8.25), Inches(1.15),
                                width=Inches(4.35))
    except Exception:
        pass
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.25), Inches(6.05),
                                  Inches(2.15), Inches(0.32))
    _fill(badge, TEAL)
    badge.line.fill.background()
    bp = badge.text_frame.paragraphs[0]
    br = bp.add_run()
    br.text = "v1 default \u00b7 v2 opt-in"
    br.font.size = Pt(9.5)
    br.font.color.rgb = WHITE
    br.font.bold = True
    return slide


def add_bullets(slide, items, x, y, w, h, size=13, color=NAVY, gap=0.06):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(6)
        p.space_before = Pt(2)
        r = p.add_run()
        r.text = ("\u2022  " if level == 0 else "\u2013  ") + text
        r.font.size = Pt(size - (1 if level else 0))
        r.font.color.rgb = color if level == 0 else GRAY
        r.font.name = "Calibri"
    return tb


def add_text(slide, text, x, y, w, h, size=12, color=NAVY, bold=False, italic=False,
             align=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return tb


def add_image(slide, name, x, y, w=None, h=None):
    path = FIG / name
    if not path.exists():
        return None
    kwargs = {}
    if w is not None:
        kwargs["width"] = Inches(w)
    if h is not None:
        kwargs["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)


def add_table(slide, data, x, y, w, h, font_size=9.5, header=True,
              right_cols=(), highlight_rows=(), shade_rows=()):
    rows = len(data)
    cols = len(data[0])
    tbl = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    col_w = w / cols
    for ci in range(cols):
        tbl.columns[ci].width = Inches(col_w)
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if ci in right_cols else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(font_size)
            r.font.name = "Calibri"
            if ri == 0 and header:
                r.font.bold = True
                r.font.color.rgb = WHITE
                _fill(cell, NAVY)
            else:
                r.font.color.rgb = NAVY
                if ri in highlight_rows:
                    _fill(cell, RGBColor(0xDD, 0xEF, 0xF2))
                elif ri in shade_rows:
                    _fill(cell, RGBColor(0xFB, 0xEA, 0xDD))
                else:
                    _fill(cell, WHITE)
    return tbl


def add_flow(slide, labels, x, y, w, h):
    n = len(labels)
    arrow_w = 0.22
    box_w = (w - arrow_w * (n - 1)) / n
    cx = x
    for i, label in enumerate(labels):
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(y),
                                    Inches(box_w), Inches(h))
        _fill(shp, LIGHT)
        shp.line.color.rgb = TEAL
        shp.line.width = Pt(1.2)
        tf = shp.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = NAVY
        r.font.name = "Calibri"
        cx += box_w
        if i < n - 1:
            arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx), Inches(y + h / 2 - 0.12),
                                        Inches(arrow_w), Inches(0.24))
            _fill(arr, TEAL)
            arr.line.fill.background()
            cx += arrow_w


def add_notes(slide, text):
    notes = slide.notes_slide.notes_text_frame
    notes.text = text


# ---------------------------------------------------------------- slides
slide_no = 0


def S(title=None, subtitle=None):
    global slide_no
    slide_no += 1
    if title is None:
        return title_slide()
    return content_slide(title, subtitle, n=slide_no)


# 1 ---------------------------------------------------------------------
title_slide()
add_notes(prs.slides[-1], "Welcome. FogSense is a classical, density-aware, multi-algorithm dehazing framework. "
          "The headline is not a black box: every decision is explainable and every number is reproducible. "
          "v1 is the recommended default; v2 is an opt-in extension that we evaluate honestly.")

# 2 ---------------------------------------------------------------------
S("Agenda", "What we will cover")
add_bullets(prs.slides[-1], [
    "1.  Problem: why fixed-strength DCP fails across haze levels",
    "2.  Atmospheric model and label-free fog-density estimation",
    "3.  FogSense architecture and adaptive strategy selection",
    "4.  Evaluation protocol, baselines and main SOTS results",
    "5.  v2 upgrade: stronger baselines and an honest negative result",
    "6.  Ablations, runtime, demo, failure cases and conclusions",
], 0.7, 1.5, 11.9, 4.8, size=15)
add_notes(prs.slides[-1], "Six parts, about twelve minutes. The key message: adapt the correction to estimated "
          "density, keep it classical and explainable, and report the v2 result honestly.")

# 3 ---------------------------------------------------------------------
S("Problem: one fixed strength cannot fit every haze level")
add_bullets(prs.slides[-1], [
    "Haze scatters and attenuates light; contrast and colour fidelity drop.",
    "Dark Channel Prior with one \u03c9 and t\u2080 over-corrects light haze and under-corrects dense haze.",
    "The correction strength and enhancement should depend on the image, not on a global constant.",
    "Goal: estimate fog density first, then adapt DCP and enhancement to that density.",
], 0.7, 1.45, 5.3, 4.6, size=13)
add_image(prs.slides[-1], "slide_problem.png", 6.15, 1.35, w=6.85)
add_notes(prs.slides[-1], "Show the two rows: light haze becomes too dark with fixed DCP, heavy haze stays hazy. "
          "FogSense adapts \u03c9, t\u2080 and the enhancement chain per class.")

# 4 ---------------------------------------------------------------------
S("Atmospheric model and transmission")
add_text(prs.slides[-1], "I(x) = J(x) t(x) + A (1 \u2212 t(x)),     t(x) = exp(\u2212\u03b2 d(x))",
         0.7, 1.45, 5.4, 0.5, size=15, bold=True, color=TEAL)
add_bullets(prs.slides[-1], [
    "J: clear scene radiance; A: atmospheric light; t: transmission; \u03b2d: optical depth.",
    "DCP estimates A from the brightest 0.1% of dark-channel pixels and t from the dark channel.",
    "Guided filter refines the noisy transmission map while preserving edges.",
    "Radiance recovery: J = (I \u2212 A) / max(t, t\u2080) + A.",
    "Baseline uses fixed \u03c9 = 0.95, t\u2080 = 0.10 for every image.",
], 0.7, 2.05, 5.4, 3.8, size=12.5)
add_image(prs.slides[-1], "transmission_maps.png", 6.35, 1.35, w=6.5, h=5.0)
add_notes(prs.slides[-1], "Explain the model in one minute. The guided filter is the difference between a usable "
          "transmission map and a noisy one; we implement the box-filter version locally.")

# 5 ---------------------------------------------------------------------
S("FogSense architecture")
add_flow(prs.slides[-1], ["Input", "Preprocess", "Density\nscore", "Selector", "DCP +\nguided filter", "Enhance"], 0.7, 2.0, 12.0, 0.85)
add_bullets(prs.slides[-1], [
    "Preprocessing: RGB float32, long side 640 px, light bilateral denoise. Shared by all methods.",
    "Density score \u2192 three-bin selector (light / moderate / heavy) \u2192 class-specific \u03c9, t\u2080, filter, CLAHE, Retinex.",
    "Baseline branch: fixed DCP for the paired comparison.",
    "Stronger baselines (v2): CAP (Zhu 2015) and a BCCR-style boundary-constraint variant.",
    "v2 opt-in branch: v2 airlight, sky/white-object handling, local density maps.",
], 0.7, 3.25, 12.0, 3.2, size=12.5)
add_notes(prs.slides[-1], "This is the 'multi-algorithm' part: the same density-aware interface can dispatch the adaptive "
          "pipeline and the fixed DCP, CAP and BCCR-style baselines. v2 is strictly opt-in.")

# 6 ---------------------------------------------------------------------
S("Fog-density estimation (label-free)")
add_bullets(prs.slides[-1], [
    "Dark-channel mean (15\u00d715 patch): higher \u2192 hazier. Weight 0.45.",
    "Shannon entropy of grayscale: lower \u2192 hazier. Weight 0.30.",
    "Global contrast (RMS + Michelson): lower \u2192 hazier. Weight 0.25.",
    "Each feature is percentile-normalised (p2\u2013p98) using calib statistics only.",
    "score = 0.45\u00b7ndc + 0.30\u00b7(1\u2212nentropy) + 0.25\u00b7(1\u2212ncontrast).",
    "Thresholds are the calib tertiles: t\u2081 = 0.3937, t\u2082 = 0.5885.",
], 0.7, 1.45, 5.5, 4.7, size=12.5)
add_image(prs.slides[-1], "score_histograms.png", 6.45, 1.5, w=6.4, h=2.15)
add_text(prs.slides[-1], "No per-image beta label is needed on SOTS; the score is validated on the synthetic beta sweep.",
         6.45, 3.85, 6.4, 0.4, size=10, italic=True, color=GRAY)
add_notes(prs.slides[-1], "Stress that this is label-free and fitted only on calib. Entropy carries most of the signal; "
          "contrast is close to neutral but kept for complementarity on real low-contrast images.")

# 7 ---------------------------------------------------------------------
S("Density validation")
add_image(prs.slides[-1], "score_vs_beta.png", 0.7, 1.45, w=7.35, h=2.7)
add_bullets(prs.slides[-1], [
    "Overall Spearman \u03c1(score, \u03b2) = 0.907.",
    "Within-scene mean \u03c1 = 0.992 (minimum 0.939 across 8 scenes).",
    "Local density is the best single signal on the sweep: 0.948 vs 0.907 global.",
    "On RTTS the ordering flips: global density tracks FADE better (0.933 vs 0.817).",
    "Calib-fitted ridge weights reach 0.969 vs FADE but are reported as an experiment only.",
], 8.2, 1.55, 4.7, 4.2, size=12)
add_notes(prs.slides[-1], "The score is monotone in beta for every scene. The RTTS result is why the shipped estimator "
          "stays global while the ridge version is documented but not shipped.")

# 8 ---------------------------------------------------------------------
S("Adaptive strategy selection")
data = [
    ["class", "\u03c9", "t\u2080", "CLAHE clip", "Retinex"],
    ["light", "0.86", "0.35", "1.5", "off"],
    ["moderate", "0.94", "0.30", "1.5", "0.35"],
    ["heavy", "0.94", "0.20", "1.5", "off"],
]
add_table(prs.slides[-1], data, 0.7, 1.45, 5.1, 1.7, font_size=11, right_cols=(1, 2, 3))
add_text(prs.slides[-1], "Parameters fitted on calib with a constrained two-stage search; correction strength increases with density.",
         0.7, 3.25, 5.1, 0.4, size=10, italic=True, color=GRAY)
add_image(prs.slides[-1], "selected_parameters.png", 6.0, 1.45, w=7.0, h=1.85)
add_bullets(prs.slides[-1], [
    "Continuous interpolation: 18.80 dB / 0.8574 SSIM (over-enhances some images).",
    "Depth-3 classifier: 18.40 dB / 0.8780 SSIM.",
    "Three-bin rule retained: simplest, interpretable, best SSIM (0.8896).",
], 6.0, 3.5, 7.0, 2.0, size=11.5)
add_notes(prs.slides[-1], "The selector comparison is an ablation, not a claim that three bins are magically optimal. "
          "They win on SSIM and remain interpretable, which matters for a course project.")

# 9 ---------------------------------------------------------------------
S("DCP dehazing and guided filtering")
add_bullets(prs.slides[-1], [
    "Dark channel on 15\u00d715 patches estimates haze thickness.",
    "Atmospheric light from the brightest 0.1% of dark-channel pixels.",
    "Transmission: t = 1 \u2212 \u03c9 \u00b7 dark(I/A), refined by the guided filter.",
    "Radiance recovery with a t\u2080 floor to limit noise amplification.",
    "Guided filter is edge-preserving; we use the local box-filter version.",
    "Ablation: removing the guided filter costs about 0.85 dB PSNR.",
], 0.7, 1.45, 5.5, 4.7, size=12.5)
add_image(prs.slides[-1], "transmission_maps.png", 6.4, 1.35, w=6.45, h=5.1)
add_notes(prs.slides[-1], "The transmission figure shows why refinement matters: raw DCP transmission is noisy and "
          "blocky; the guided filter recovers edges and avoids halos.")

# 10 --------------------------------------------------------------------
S("Enhancement stack")
add_bullets(prs.slides[-1], [
    "Retinex: multi-scale illumination correction on luminance (moderate class only).",
    "CLAHE: adaptive histogram equalisation on L in LAB, blended at strength 0.4.",
    "Gray-world: per-channel gains remove residual colour cast.",
    "Each step is individually switchable for ablation.",
    "Light/heavy disable Retinex; heavy keeps CLAHE for detail.",
], 0.7, 1.45, 5.3, 4.6, size=12.5)
add_image(prs.slides[-1], "before_after_by_class.png", 6.2, 1.35, w=6.7, h=4.55)
add_notes(prs.slides[-1], "Enhancement is where the SSIM gain comes from. CLAHE alone reaches 18.18 dB / 0.8850; adding "
          "Retinex raises PSNR to 19.30 but lowers SSIM to 0.8840.")

# 11 --------------------------------------------------------------------
S("v2 opt-in upgrades")
add_bullets(prs.slides[-1], [
    "Airlight v2: brightest 0.5% of dark channel, reject small components and high-gradient pixels; fallback to v1.",
    "Sky / white-object mask: bright + smooth + connected; floors transmission at 0.25 and blends 25% toward input.",
    "Local density map: local dark channel, rank entropy and contrast; guided-filter smoothed.",
    "Per-pixel \u03c9(x), t\u2080(x) interpolated from class centroids; enhancement blended by local density.",
    "Quad-tree airlight search rejected on calib (\u22120.88 dB), so it is not used.",
], 0.7, 1.45, 5.5, 4.7, size=12.5)
add_image(prs.slides[-1], "v2_density_maps.png", 6.4, 1.35, w=6.45, h=4.6)
add_text(prs.slides[-1], "The green mask shows the conservative sky/white-object region; it can false-positive on bright indoor areas.",
         6.4, 6.0, 6.45, 0.35, size=9, italic=True, color=GRAY)
add_notes(prs.slides[-1], "v2 is deliberately conservative: airlight changes PSNR by only +0.11 dB on calib. The sky stage is "
          "the only step that improves both PSNR and SSIM; the local maps are the risky part.")

# 12 --------------------------------------------------------------------
S("Evaluation protocol and baselines")
add_bullets(prs.slides[-1], [
    "Data: calib 299 pairs, SOTS test 243 pairs, RTTS 72 images, density validation 64, selector set 180.",
    "Split by unique clear scene; preprocessing identical for every method.",
    "Frozen tags: v0.1-frozen-params, v1.0, v2.0-frozen-params, v2.0. No test-set tuning.",
    "Baselines: fixed DCP, CAP (Zhu 2015), BCCR-style (Meng 2013, simplified).",
    "Metrics: PSNR/SSIM, CIEDE2000, NIQE-style and FADE-style (proxies), entropy, contrast, edge, runtime.",
    "All v2 knobs fitted on calib only; v1 tables remain unchanged.",
], 0.7, 1.45, 12.0, 4.8, size=12.5)
add_notes(prs.slides[-1], "Emphasise the freeze discipline: parameters are written before the test run and the repository is "
          "tagged. NIQE/FADE are locally fitted or fixed-weight proxies, not official models.")

# 13 --------------------------------------------------------------------
S("Main results: SOTS test (243 pairs)")
data = [
    ["method", "PSNR (dB)", "SSIM", "CIEDE2000", "ms/image"],
    ["DCP", "15.71", "0.8263", "12.64", "59"],
    ["CAP", "18.09", "0.7833", "9.84", "38"],
    ["BCCR-style", "16.02", "0.8309", "12.36", "83"],
    ["FogSense v1", "18.78", "0.8896", "9.50", "192"],
    ["FogSense v2", "18.45", "0.8873", "9.38", "562"],
]
add_table(prs.slides[-1], data, 0.7, 1.45, 7.2, 2.0, font_size=10.5,
          right_cols=(1, 2, 3, 4), highlight_rows=(4,), shade_rows=(5,))
add_bullets(prs.slides[-1], [
    "FogSense v1: +3.07 dB PSNR and +0.0633 SSIM over fixed DCP.",
    "Win rate v1 vs DCP: PSNR 96.3%, SSIM 90.1%; gains in every haze class.",
    "CAP is a much stronger PSNR baseline than DCP (18.09 vs 15.71) but lower SSIM.",
    "BCCR-style is the most 'natural' (best NIQE) but under-enhances detail.",
    "v1 also leads on real RTTS entropy gain (0.4165) and edge preservation (1.6733).",
], 8.1, 1.45, 4.9, 4.6, size=12)
add_notes(prs.slides[-1], "Read the table slowly. v1 is best on PSNR and SSIM; v2 slightly better on colour (CIEDE). "
          "CAP narrows the PSNR advantage but not the structural/colour advantage.")

# 14 --------------------------------------------------------------------
S("Honest v2 result: better indoors, worse on average")
add_image(prs.slides[-1], "slide_v2_delta.png", 0.7, 1.35, w=7.55, h=3.45)
add_bullets(prs.slides[-1], [
    "v2 vs v1: PSNR \u22120.32 dB (p = 0.0002), SSIM \u22120.0023 (p = 0.44).",
    "Indoor (n = 22): +0.75 dB PSNR, +0.0118 SSIM.",
    "Outdoor (n = 221): \u22120.43 dB PSNR, \u22120.0037 SSIM.",
    "Sky stage alone: +0.64 dB / +0.0037 SSIM over v1.",
    "Local maps cost \u22120.87 dB; enhancement blend recovers SSIM only.",
    "Brightness diagnostic: after mean/std matching, v2 is +0.37 dB vs v1.",
    "Conclusion: v1 remains the default; v2 stays opt-in.",
], 8.45, 1.45, 4.5, 4.8, size=11.5)
add_notes(prs.slides[-1], "This is the honest result and a strength of the project. The sky stage is genuinely useful; "
          "the local omega/t0 maps are not, as configured. We keep the frozen config rather than tuning on the test set.")

# 15 --------------------------------------------------------------------
S("Ablations: what actually helps")
add_image(prs.slides[-1], "slide_ablations.png", 0.7, 1.35, w=12.0, h=3.95)
add_bullets(prs.slides[-1], [
    "Enhancement drives SSIM: no enhancement 0.8726 \u2192 full pipeline 0.8896.",
    "Retinex raises PSNR to 19.30 dB but lowers SSIM to 0.8840; gray-world gives best SSIM and lowest colour cast.",
    "Three-bin selector beats continuous (0.8574 SSIM) and the depth-3 classifier (0.8780).",
    "Density: dark-channel only 18.48 / 0.8831; adding entropy reaches 18.84 / 0.8901.",
    "v2: only the sky step improves both metrics; local maps lose PSNR.",
], 0.7, 5.45, 12.0, 1.5, size=11.5)
add_notes(prs.slides[-1], "Use the ablation figure to justify the final design choices. The small SSIM difference between "
          "entropy-only and full density is within noise, so we do not over-claim it.")

# 16 --------------------------------------------------------------------
S("Qualitative results")
add_image(prs.slides[-1], "slide_qualitative.png", 0.7, 1.35, w=12.0, h=3.95)
add_bullets(prs.slides[-1], [
    "v1 recovers contrast and detail across haze levels; v2 helps indoor bright regions.",
    "CAP is strong on PSNR but has lower SSIM; BCCR-style is most natural but under-enhances.",
    "v2 can over- or under-correct outdoors, matching the quantitative outdoor loss.",
], 0.7, 5.45, 12.0, 1.5, size=11.5)
add_notes(prs.slides[-1], "Point at the sky regions: v2's conservative mask helps some images and hurts others. "
          "The full four-row figure is in the report.")

# 17 --------------------------------------------------------------------
S("Runtime and live demo")
add_image(prs.slides[-1], "slide_runtime.png", 0.7, 1.35, w=6.9, h=3.7)
add_image(prs.slides[-1], "app_demo.png", 7.85, 1.35, w=5.0, h=3.15)
add_bullets(prs.slides[-1], [
    "Fixed DCP 59 ms; FogSense v1 192 ms; FogSense v2 562 ms (~2.9\u00d7 v1).",
    "Local density map is the main extra cost; airlight/sky are cheap.",
    "Live demo: Streamlit app with density map, transmission map and v2 toggles.",
], 0.7, 5.2, 6.9, 1.6, size=11.5)
add_text(prs.slides[-1], "python -m streamlit run app/app.py", 7.85, 4.6, 5.0, 0.35,
         size=10.5, color=TEAL, bold=True)
add_notes(prs.slides[-1], "The screenshot is from the real app. Mention that the live timings vary with CPU load; the "
          "benchmark numbers are from the committed evaluation run.")

# 18 --------------------------------------------------------------------
S("Failure cases and limitations")
add_bullets(prs.slides[-1], [
    "Sky and white objects violate the dark-channel assumption \u2192 darkening or colour shift.",
    "Night haze has weak colour statistics; the density score is less reliable.",
    "Noise amplification in dense haze: the t\u2080 floor trades detail for noise.",
    "CLAHE/Retinex can over-sharpen texture and amplify artefacts.",
    "v2: indoor gain from only 22 pairs; outdoor loss dominates the average.",
    "NIQE-style and FADE-style are proxies, not official learned models.",
    "v2 is ~2.9\u00d7 slower than v1 and ~9\u00d7 slower than fixed DCP on CPU.",
], 0.7, 1.45, 5.6, 5.0, size=12)
add_image(prs.slides[-1], "failure_cases.png", 6.5, 1.35, w=6.4, h=4.65)
add_notes(prs.slides[-1], "Be explicit about limitations; this is what makes the evaluation credible. The sky-mask "
          "false positives are visible in the v2 density-map figure.")

# 19 --------------------------------------------------------------------
S("Conclusions")
add_bullets(prs.slides[-1], [
    "FogSense is a classical, density-aware, multi-algorithm framework with an explainable decision chain.",
    "v1 reaches 18.78 dB / 0.8896 SSIM vs fixed DCP 15.71 dB / 0.8263 SSIM on 243 SOTS pairs.",
    "Gains in every haze class; 96.3% PSNR and 90.1% SSIM win rates.",
    "CAP is a strong baseline; the adaptive advantage is structural and colour, not only PSNR.",
    "v2 is reported honestly: better indoors and on CIEDE2000, but an aggregate PSNR loss; v1 stays default.",
    "Future work: better airlight/sky handling, adaptive t\u2080, faster local features, real-world validation.",
], 0.7, 1.45, 12.0, 4.9, size=12.5)
add_notes(prs.slides[-1], "Close on the honest v2 result: we learned that the sky stage is the valuable part and the local "
          "maps need rethinking. That is a real, defensible conclusion.")

# 20 --------------------------------------------------------------------
S("Questions?")
add_text(prs.slides[-1], "Thank you", 0.7, 1.9, 12.0, 0.8, size=40, bold=True, color=TEAL)
add_bullets(prs.slides[-1], [
    "Repository: github.com/AtulTejaswi/adaptive-fog-density-dehazing",
    "Demo: python -m streamlit run app/app.py",
    "Report: report/report.md  \u00b7  Slides: report/FogSense_slides.pptx",
    "All numbers are generated from files under results/ and are reproducible.",
], 0.7, 3.0, 12.0, 2.5, size=13)
add_notes(prs.slides[-1], "Invite questions on the density estimator, the v2 negative result, or the baseline fairness audit.")

# 21 appendix ------------------------------------------------------------
S("Appendix A: metric definitions", "Backup slide")
add_bullets(prs.slides[-1], [
    "PSNR / SSIM: full-reference against the clear SOTS ground truth.",
    "CIEDE2000: mean colour difference in Lab (lower is better).",
    "NIQE-style: MSCN statistics vs a Gaussian fitted on 299 calib clear images (lower is better).",
    "FADE-style: Choi et al. feature set with fixed documented weights (lower = clearer).",
    "No-reference: entropy gain, RMS contrast improvement, Sobel edge-preservation ratio, colour cast.",
    "NIQE/FADE are proxies: relative trends only, not absolute scores.",
], 0.7, 1.45, 12.0, 4.8, size=12)
add_notes(prs.slides[-1], "Use if asked about metrics. Be clear that NIQE/FADE are locally fitted or fixed-weight proxies.")

# 22 appendix ------------------------------------------------------------
S("Appendix B: reproducibility", "Backup slide")
add_bullets(prs.slides[-1], [
    "v1 pipeline: .\\run_all.ps1 (or the individual scripts in the README).",
    "v2 pipeline: .\\run_all_v2.ps1 --no-download (about 35\u201345 minutes).",
    "Tests: python -m pytest  (83 tests, including byte-identical v2-off compatibility).",
    "Tags: v0.1-frozen-params, v1.0, v2.0-frozen-params, v2.0.",
    "With every v2 flag off, output is byte-identical to v1.0.",
], 0.7, 1.45, 12.0, 4.8, size=12)
add_notes(prs.slides[-1], "Use if asked about reproducibility or the freeze discipline.")

prs.save(OUT)
print(f"saved {OUT}  ({len(prs.slides)} slides)")
